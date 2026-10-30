#!/usr/bin/env python3
"""Generate docs/report.docx from the Markdown report source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD = ROOT / "report.md"
OUT = ROOT / "report.docx"
BODY_FONT_CN = "宋体"
HEAD_FONT_CN = "黑体"
LATIN_FONT = "Times New Roman"
CODE_FONT = "Consolas"
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT_BLUE = "EAF2FB"
CODE_BG = "F5F7FA"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]*`)")


def set_run_font(run, size=12, bold=None, color=None, latin=LATIN_FONT,
                 east_asia=BODY_FONT_CN):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_inline(paragraph, text, size=12, bold=False, color=None):
    text = text.replace("\\|", "|")
    text = re.sub(r"<(https?://[^>]+)>", r"\1", text)
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, bold=bold, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size - 1.0, bold=bold, color=color,
                         latin=CODE_FONT)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=bold, color=color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.5,
                          first_indent=None, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    if first_indent is None:
        fmt.first_line_indent = Pt(0)
    else:
        fmt.first_line_indent = Cm(first_indent)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("left", left),
                        ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, pct=5000):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(pct))
    tbl_w.set(qn("w:type"), "pct")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_borders(table, color="B8C2CC", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def add_toc_styles(doc):
    for name, indent in (("TOC 1", 0.0), ("TOC 2", 0.55)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = LATIN_FONT
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fmt = style.paragraph_format
        fmt.left_indent = Cm(indent)
        fmt.first_line_indent = Cm(0)
        fmt.space_before = Pt(1)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.1
        fmt.tab_stops.add_tab_stop(Cm(15.8), WD_TAB_ALIGNMENT.RIGHT,
                                   WD_TAB_LEADER.DOTS)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), BODY_FONT_CN)
        rfonts.set(qn("w:ascii"), LATIN_FONT)
        rfonts.set(qn("w:hAnsi"), LATIN_FONT)


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目  录")
    set_run_font(run, size=16, bold=True, latin="SimHei",
                 east_asia=HEAD_FONT_CN)
    set_paragraph_spacing(p, before=0, after=12, line=1.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=6, line=1.25)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Word 中自动更新（也可右键“更新域”）"
    run._r.append(placeholder)
    run._r.append(fld_end)


def add_header_footer(section):
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("计算机网络实验报告")
    set_run_font(run, size=9, color=GRAY, latin=LATIN_FONT)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "B8C2CC")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    add_field(fp, "PAGE")
    run = fp.add_run(" 页 / 共 ")
    set_run_font(run, size=9, color=GRAY)
    add_field(fp, "NUMPAGES")
    run = fp.add_run(" 页")
    set_run_font(run, size=9, color=GRAY)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT_CN)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)

    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6 if level == 1 else 4)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True
        srpr = style.element.get_or_add_rPr()
        srfonts = srpr.find(qn("w:rFonts"))
        if srfonts is None:
            srfonts = OxmlElement("w:rFonts")
            srpr.append(srfonts)
        srfonts.set(qn("w:eastAsia"), HEAD_FONT_CN)
        srfonts.set(qn("w:ascii"), LATIN_FONT)
        srfonts.set(qn("w:hAnsi"), LATIN_FONT)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(p, text, size=16 if level == 1 else 14, bold=True)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=0, line=1.5,
                          first_indent=0.74 if indent else None)
    add_inline(p, text, size=12)
    return p


def add_bullet(doc, text, marker="- "):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.left_indent = Cm(0.9)
    fmt.first_line_indent = Cm(-0.55)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.5
    run = p.add_run(marker)
    set_run_font(run, size=12)
    add_inline(p, text, size=12)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=8, line=1.0)
    add_inline(p, text, size=9.5, color=GRAY)
    return p


IMAGE_WIDTHS = {
    "architecture.png": 6.2,
    "protocol_sequence.png": 6.2,
    "server_main_flow.png": 4.8,
    "server_worker_flow.png": 5.9,
    "client_main_flow.png": 5.0,
    "client_receiver_flow.png": 4.15,
}


def add_image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=0, line=1.0, keep_next=True)
    run = p.add_run()
    width = IMAGE_WIDTHS.get(path.name, 6.0)
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def split_table_row(line):
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    cells = re.split(r"(?<!\\)\|", raw)
    return [cell.strip().replace("\\|", "|") for cell in cells]


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", cell.strip()) for cell in cells)


def add_table(doc, rows):
    if len(rows) < 2:
        return
    header = rows[0]
    body = [row for row in rows[1:] if not is_separator_row(row)]
    separators = rows[1] if len(rows) > 1 and is_separator_row(rows[1]) else []
    aligns = [WD_ALIGN_PARAGRAPH.LEFT]
    for c in range(1, len(header)):
        values = [row[c] for row in body if c < len(row)]
        aligns.append(WD_ALIGN_PARAGRAPH.CENTER
                      if values and max(len(v) for v in values) <= 12
                      else WD_ALIGN_PARAGRAPH.LEFT)

    table = doc.add_table(rows=len(body) + 1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_width(table)
    set_table_borders(table)

    font_size = 9.5 if len(header) <= 4 else 9.0
    for col, text in enumerate(header):
        cell = table.rows[0].cells[col]
        shade_cell(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, line=1.0)
        add_inline(p, text, size=font_size, bold=True)
        p.paragraph_format.keep_with_next = True
    set_repeat_header(table.rows[0])
    set_no_split(table.rows[0])

    for r, row in enumerate(body, start=1):
        for c in range(len(header)):
            cell = table.rows[r].cells[c]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = aligns[c] if c > 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, line=1.0)
            text = row[c] if c < len(row) else ""
            add_inline(p, text, size=font_size)
        set_no_split(table.rows[r])
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=0, line=1.0)


def add_code_block(doc, lines):
    max_len = max((len(line) for line in lines), default=0)
    if max_len > 100:
        font_size = 7.0
    elif max_len > 85:
        font_size = 7.6
    elif max_len > 72:
        font_size = 8.1
    else:
        font_size = 8.6

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table, color="D5DBE3", size="4")
    cell = table.rows[0].cells[0]
    shade_cell(cell, CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        fmt.first_line_indent = Pt(0)
        run = p.add_run(line)
        set_run_font(run, size=font_size, latin=CODE_FONT,
                     east_asia=BODY_FONT_CN)
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=0, line=1.0)


def add_cover(doc):
    for _ in range(2):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, line=1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("计算机网络实验报告")
    set_run_font(run, size=20, bold=True, latin="SimHei",
                 east_asia=HEAD_FONT_CN)
    set_paragraph_spacing(p, before=0, after=24, line=1.25)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("基于 TCP 流式套接字的多用户聊天程序")
    set_run_font(run, size=24, bold=True, latin="SimHei",
                 east_asia=HEAD_FONT_CN)
    set_paragraph_spacing(p, before=12, after=6, line=1.3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("——协议设计、多线程实现与丢包观察")
    set_run_font(run, size=14, bold=False, latin=LATIN_FONT,
                 east_asia=BODY_FONT_CN)
    set_paragraph_spacing(p, before=0, after=50, line=1.25)

    info = [
        ("课程名称", "计算机网络"),
        ("实验名称", "流式 Socket 聊天程序"),
        ("姓    名", "杨斯源"),
        ("学    号", "2412796"),
        ("日    期", "2026 年 10 月 28 日"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    for row_idx, (label, value) in enumerate(info):
        row = table.rows[row_idx]
        row.height = Cm(0.95)
        c0, c1 = row.cells
        c0.width = Cm(3.4)
        c1.width = Cm(7.8)
        for cell, text, align in ((c0, label, WD_ALIGN_PARAGRAPH.RIGHT),
                                  (c1, value, WD_ALIGN_PARAGRAPH.LEFT)):
            p = cell.paragraphs[0]
            p.alignment = align
            set_paragraph_spacing(p, line=1.5)
            run = p.add_run(text)
            set_run_font(run, size=14, latin=LATIN_FONT,
                         east_asia=BODY_FONT_CN)

    for _ in range(3):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, line=1.0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 年 10 月")
    set_run_font(run, size=12, color=GRAY)


def extract_abstract(md_text):
    abstract = []
    for line in md_text.splitlines():
        if line.startswith("# 1 "):
            break
        if line.startswith("# 摘要"):
            continue
        if line.startswith("<!--"):
            continue
        if line.strip():
            abstract.append(line.strip())
    return abstract


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.left_indent = Cm(0.74)
    fmt.first_line_indent = Cm(-0.74)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.3
    add_inline(p, text, size=10.5)
    return p


def parse_report_body(doc, md_text):
    lines = md_text.splitlines()
    try:
        start = next(i for i, line in enumerate(lines)
                     if line.startswith("# 1 "))
    except StopIteration:
        raise RuntimeError("report.md: missing '# 1 ' heading")
    lines = lines[start:]
    i = 0
    skip_caption = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if skip_caption and stripped.startswith("**图 "):
            skip_caption = False
            i += 1
            continue
        skip_caption = False
        if stripped.startswith("# "):
            add_heading(doc, stripped[2:].strip(), 1)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), 2)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), 3)
            i += 1
            continue
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            if i < len(lines):
                i += 1
            add_code_block(doc, code)
            continue
        if stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                caption, rel = match.groups()
                path = ROOT / rel
                if path.exists():
                    add_image(doc, path, caption)
                    skip_caption = True
                else:
                    add_body(doc, f"[缺少图片：{rel}]", indent=False)
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:].strip(), marker="• ")
            i += 1
            continue
        num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num:
            add_bullet(doc, num.group(2), marker=f"{num.group(1)}. ")
            i += 1
            continue
        if stripped.startswith("**") and stripped.endswith("**") and \
                stripped[2:-2].startswith("图 "):
            add_caption(doc, stripped[2:-2])
            i += 1
            continue
        if re.match(r"^\[\d+\]\s+", stripped):
            add_reference(doc, stripped)
            i += 1
            continue
        add_body(doc, stripped)
        i += 1


def main():
    md_text = MD.read_text(encoding="utf-8")
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.2)
    add_header_footer(section)

    add_cover(doc)
    doc.add_page_break()

    abstract = extract_abstract(md_text)
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("摘  要")
    set_run_font(run, size=16, bold=True, latin="SimHei",
                 east_asia=HEAD_FONT_CN)
    set_paragraph_spacing(p, before=0, after=10, line=1.25)

    keywords = None
    for text in abstract:
        if text.startswith("**关键词：**"):
            keywords = text
        else:
            add_body(doc, text)
    if keywords:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=8, after=0, line=1.5)
        add_inline(p, keywords, size=12)

    doc.add_page_break()
    add_toc_styles(doc)
    add_toc(doc)
    doc.add_page_break()

    parse_report_body(doc, md_text)

    core = doc.core_properties
    core.title = "基于 TCP 流式套接字的多用户聊天程序"
    core.subject = "计算机网络实验报告"
    core.author = "杨斯源"
    core.comments = "CHAT/1, TCP, pthread, UTF-8"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
